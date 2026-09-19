"""Build the client presentation from real, verified renders.

PDF via ReportLab and an editable PPTX via python-pptx -- both free, both local.
Nothing is fabricated: a concept with no render on disk is reported as missing
rather than represented by a placeholder.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (Image as RLImage, PageBreak, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

from backend.app.core.config import DELIVERABLES, ROOT
from backend.app.database.db import init_db
from backend.app.versions import store

INK = colors.HexColor("#14181D")
MARIGOLD = colors.HexColor("#E8890C")
MUTE = colors.HexColor("#5A6672")

ORDER = ["Concept A Balanced Split", "Concept B S River Split",
         "Concept C Organic Asymmetric"]

SHOT_LABELS = {
    "01_complete_flower": "Assembled flower",
    "02_separated_components": "The two components, separated",
    "03_top_view": "Plan view",
    "04_three_quarter_view": "Three-quarter view",
    "05_split_closeup": "Detail of the division",
    "06_side_view": "Side elevation",
}


def styles():
    s = getSampleStyleSheet()
    s.add(ParagraphStyle("Cover", parent=s["Title"], fontSize=34, leading=40,
                         textColor=INK, spaceAfter=6))
    s.add(ParagraphStyle("Sub", parent=s["Normal"], fontSize=13, leading=18,
                         textColor=MUTE))
    s.add(ParagraphStyle("H", parent=s["Heading1"], fontSize=20, leading=25,
                         textColor=INK, spaceAfter=4))
    s.add(ParagraphStyle("Body", parent=s["Normal"], fontSize=10.5, leading=15,
                         textColor=INK))
    s.add(ParagraphStyle("Small", parent=s["Normal"], fontSize=8.5, leading=12,
                         textColor=MUTE))
    s.add(ParagraphStyle("Caption", parent=s["Normal"], fontSize=8, leading=10,
                         textColor=MUTE, alignment=1))
    return s


def collect() -> list[dict]:
    init_db()
    project = store.ensure_default_project()
    out = []
    for concept in store.list_concepts(project["id"], include_archived=False):
        if concept["name"] not in ORDER:
            continue
        head = concept.get("head")
        if not head:
            continue
        assets = head.get("assets") or {}
        shots = {k: ROOT / v for k, v in assets.items()
                 if k in SHOT_LABELS and (ROOT / v).is_file()}
        out.append({"concept": concept, "version": head, "shots": shots,
                    "validation": head.get("validation") or {},
                    "stats": head.get("stats") or {}})
    out.sort(key=lambda d: ORDER.index(d["concept"]["name"]))
    return out


_JPEG_CACHE: dict[Path, Path] = {}


def _as_jpeg(path: Path) -> Path:
    """Re-encode a render as JPEG for embedding.

    ReportLab embeds PNGs losslessly, which makes a deck of eighteen renders
    over 20 MB -- too big to email, for no visible gain at print size.
    """
    if path in _JPEG_CACHE:
        return _JPEG_CACHE[path]
    from PIL import Image as PILImage
    out = DELIVERABLES / "_pdf_cache" / (path.stem + ".jpg")
    out.parent.mkdir(parents=True, exist_ok=True)
    with PILImage.open(path) as im:
        im.convert("RGB").save(out, "JPEG", quality=88, optimize=True,
                               progressive=True)
    _JPEG_CACHE[path] = out
    return out


def fit(path: Path, max_w: float, max_h: float) -> RLImage:
    from PIL import Image as PILImage
    with PILImage.open(path) as im:
        w, h = im.size
    scale = min(max_w / w, max_h / h)
    return RLImage(str(_as_jpeg(Path(path))), width=w * scale, height=h * scale)


def build_pdf(data: list[dict], out_path: Path) -> dict:
    s = styles()
    doc = SimpleDocTemplate(str(out_path), pagesize=landscape(A4),
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=13 * mm, bottomMargin=13 * mm,
                            title="AFRi Marigold Accessory - Initial Concepts",
                            author="AFRi Studio")
    W = doc.width
    story = []
    missing = []

    # ---- cover ----------------------------------------------------------
    story += [Spacer(1, 40 * mm),
              Paragraph("AFRi Marigold Accessory", s["Cover"]),
              Paragraph("Initial Concepts &mdash; Stage One", s["Sub"]),
              Spacer(1, 8 * mm),
              Paragraph(
                  "Three provisional approaches to dividing a single marigold "
                  "into two separate pieces that reassemble into one complete "
                  "flower.", s["Body"]),
              Spacer(1, 14 * mm),
              Paragraph(
                  "<b>These concepts are provisional.</b> AFRi's reference "
                  "material was not available when they were produced, so the "
                  "colours, proportions and construction shown here are studio "
                  "assumptions, not brand specifications. No claim is made that "
                  "they match AFRi's existing design language.", s["Small"]),
              PageBreak()]

    # ---- objective ------------------------------------------------------
    story += [Paragraph("The design objective", s["H"]),
              Paragraph(
                  "The accessory is a single marigold made in two separate "
                  "pieces. Apart, each piece is a complete object in its own "
                  "right. Brought together, they reconstruct one whole flower, "
                  "with the line of the division reading as a deliberate part "
                  "of the design rather than a seam.", s["Body"]),
              Spacer(1, 5 * mm),
              Paragraph(
                  "All three concepts are cut from the <b>same master flower</b>. "
                  "Only the dividing line differs between them, so the "
                  "comparison is a fair one: what changes from concept to "
                  "concept is the division itself, not the flower.", s["Body"]),
              Spacer(1, 5 * mm)]

    rows = [["", "Confirmed requirement"],
            ["1", "The accessory is a marigold."],
            ["2", "It is made in two separate pieces."],
            ["3", "Together the pieces form one complete flower."],
            ["4", "Several ways of dividing the flower are to be explored."],
            ["5", "An S-shaped, river-inspired division is a priority."],
            ["6", "It will sit on a hat brim (a later stage, not shown here)."]]
    t = Table(rows, colWidths=[12 * mm, W - 12 * mm])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9.5),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), INK),
        ("TEXTCOLOR", (0, 1), (-1, -1), INK),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("LINEBELOW", (0, 0), (-1, 0), 1, MARIGOLD)]))
    story += [t, PageBreak()]

    # ---- one spread per concept ----------------------------------------
    for entry in data:
        c, v = entry["concept"], entry["version"]
        shots = entry["shots"]
        story.append(Paragraph(c["name"], s["H"]))
        story.append(Paragraph(c["description"], s["Body"]))
        story.append(Spacer(1, 4 * mm))

        main = ["01_complete_flower", "02_separated_components", "05_split_closeup"]
        cells, caps = [], []
        for key in main:
            if key in shots:
                cells.append(fit(shots[key], W / 3 - 8 * mm, 88 * mm))
                caps.append(Paragraph(SHOT_LABELS[key], s["Caption"]))
            else:
                missing.append(f"{c['name']}: {key}")
                cells.append(Paragraph("(not rendered)", s["Caption"]))
                caps.append(Paragraph(SHOT_LABELS[key], s["Caption"]))
        if cells:
            g = Table([cells, caps], colWidths=[W / 3] * 3)
            g.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"),
                                   ("VALIGN", (0, 0), (-1, 0), "MIDDLE"),
                                   ("TOPPADDING", (0, 1), (-1, 1), 3)]))
            story.append(g)

        st = entry["stats"]
        val = entry["validation"]
        meas = (val.get("measurements") or {})
        facts = [
            f"Version {v['number']}",
            f"{st.get('petal_count', '?')} petals in {st.get('layer_count', '?')} rows",
            f"{st.get('diameter_mm', '?')} mm across, {st.get('height_mm', '?')} mm deep",
            f"piece balance {meas.get('volume_balance', '?')}",
            f"validation {val.get('passed', '?')}/{val.get('total', '?')}",
        ]
        story += [Spacer(1, 3 * mm),
                  Paragraph(" &nbsp;·&nbsp; ".join(facts), s["Small"]),
                  PageBreak()]

        # secondary views
        rest = [k for k in ("03_top_view", "04_three_quarter_view", "06_side_view")
                if k in shots]
        if rest:
            story.append(Paragraph(f"{c['name']} &mdash; further views", s["H"]))
            story.append(Spacer(1, 3 * mm))
            cells = [fit(shots[k], W / len(rest) - 8 * mm, 95 * mm) for k in rest]
            caps = [Paragraph(SHOT_LABELS[k], s["Caption"]) for k in rest]
            g = Table([cells, caps], colWidths=[W / len(rest)] * len(rest))
            g.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"),
                                   ("TOPPADDING", (0, 1), (-1, 1), 3)]))
            story += [g, PageBreak()]

    # ---- comparison -----------------------------------------------------
    comp = [e for e in data if "01_complete_flower" in e["shots"]]
    if comp:
        story.append(Paragraph("The three concepts side by side", s["H"]))
        story.append(Paragraph(
            "Rendered at matched scale, framing and lighting, from the same "
            "master flower. Assembled above, separated below.", s["Small"]))
        story.append(Spacer(1, 2 * mm))
        # Sized so the heading and both rows fit one landscape page: overflowing
        # leaves the heading stranded alone on the page before.
        cw = W / len(comp) - 8 * mm
        cells = [fit(e["shots"]["01_complete_flower"], cw, 68 * mm) for e in comp]
        caps = [Paragraph(e["concept"]["name"], s["Caption"]) for e in comp]
        sep = [fit(e["shots"]["02_separated_components"], cw, 62 * mm)
               if "02_separated_components" in e["shots"] else
               Paragraph("(not rendered)", s["Caption"]) for e in comp]
        g = Table([cells, caps, sep], colWidths=[W / len(comp)] * len(comp))
        g.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"),
                               ("TOPPADDING", (0, 1), (-1, 1), 3)]))
        story += [g, PageBreak()]

    # ---- next steps -----------------------------------------------------
    story += [
        Paragraph("Next steps", s["H"]),
        Paragraph(
            "Please select a direction, or tell us which elements of each you "
            "would like combined. Every parameter shown is adjustable: petal "
            "count and density, depth of relief, the shape and position of the "
            "dividing line, and the finish.", s["Body"]),
        Spacer(1, 4 * mm),
        Paragraph("<b>What we would like from you</b>", s["Body"]),
        Paragraph(
            "1. A preferred concept, or a combination.<br/>"
            "2. AFRi's reference material, so the provisional assumptions can be "
            "replaced with your actual design language.<br/>"
            "3. The intended finished size and the material you have in mind.<br/>"
            "4. Confirmation to proceed to the hat stage, where the approved "
            "flower is placed on the brim and the attachment is worked out.",
            s["Body"]),
        Spacer(1, 8 * mm),
        Paragraph(
            "<b>Scope and caveats.</b> These are digital design concepts. The "
            "hat is not modelled and no attachment mechanism has been designed "
            "&mdash; that is deliberately a later stage. Manufacturing "
            "feasibility has not been assessed: wall thickness, tolerances and "
            "process constraints all remain open. The colours are provisional "
            "and are not AFRi brand colours.", s["Small"]),
    ]

    doc.build(story)
    return {"pdf": str(out_path), "missing": missing,
            "concepts": len(data),
            "bytes": out_path.stat().st_size if out_path.is_file() else 0}


def build_pptx(data: list[dict], out_path: Path) -> dict:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Emu, Inches, Pt
    except ImportError as exc:
        return {"pptx": None, "error": f"python-pptx unavailable: {exc}"}

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, s, size=18, bold=False, colour=(20, 24, 29)):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = s
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor(*colour)
        return box

    s0 = prs.slides.add_slide(blank)
    text(s0, 0.9, 2.4, 11.5, 1.2, "AFRi Marigold Accessory", 44, True)
    text(s0, 0.9, 3.5, 11.5, 0.8, "Initial Concepts - Stage One", 20,
         colour=(90, 102, 114))
    text(s0, 0.9, 5.2, 11.5, 1.6,
         "Provisional concepts. AFRi reference material was not available, so "
         "colours, proportions and construction are studio assumptions.",
         12, colour=(90, 102, 114))

    for entry in data:
        c = entry["concept"]
        shots = entry["shots"]
        sl = prs.slides.add_slide(blank)
        text(sl, 0.6, 0.35, 12, 0.7, c["name"], 28, True)
        text(sl, 0.6, 1.05, 12, 0.6, c["description"], 12, colour=(90, 102, 114))
        keys = [k for k in ("01_complete_flower", "02_separated_components",
                            "05_split_closeup") if k in shots]
        for i, k in enumerate(keys):
            sl.shapes.add_picture(str(shots[k]), Inches(0.6 + i * 4.15),
                                  Inches(1.9), height=Inches(4.3))
            text(sl, 0.6 + i * 4.15, 6.3, 4.0, 0.4, SHOT_LABELS[k], 10,
                 colour=(90, 102, 114))

    comp = [e for e in data if "01_complete_flower" in e["shots"]]
    if comp:
        sl = prs.slides.add_slide(blank)
        text(sl, 0.6, 0.35, 12, 0.7, "The three concepts side by side", 28, True)
        for i, e in enumerate(comp):
            sl.shapes.add_picture(str(e["shots"]["01_complete_flower"]),
                                  Inches(0.6 + i * 4.15), Inches(1.6),
                                  height=Inches(4.6))
            text(sl, 0.6 + i * 4.15, 6.4, 4.0, 0.4, e["concept"]["name"], 11,
                 colour=(90, 102, 114))

    sl = prs.slides.add_slide(blank)
    text(sl, 0.6, 0.5, 12, 0.7, "Next steps", 28, True)
    text(sl, 0.6, 1.5, 12, 3.2,
         "1. Select a preferred concept, or a combination.\n"
         "2. Share AFRi reference material so the provisional assumptions can "
         "be replaced.\n"
         "3. Confirm the intended finished size and material.\n"
         "4. Authorise the hat stage: placing the approved flower on the brim "
         "and designing the attachment.", 15)
    text(sl, 0.6, 5.6, 12, 1.4,
         "The hat is not modelled and no attachment has been designed. "
         "Manufacturing feasibility has not been assessed.", 11,
         colour=(90, 102, 114))

    prs.save(str(out_path))
    return {"pptx": str(out_path),
            "bytes": out_path.stat().st_size if out_path.is_file() else 0}


def main():
    DELIVERABLES.mkdir(parents=True, exist_ok=True)
    data = collect()
    if not data:
        print(json.dumps({"ok": False,
                          "error": "no concepts with renders were found"}))
        return 1
    pdf = build_pdf(data, DELIVERABLES / "AFRi_Marigold_Initial_Concepts.pdf")
    import shutil as _shutil
    _shutil.rmtree(DELIVERABLES / "_pdf_cache", ignore_errors=True)
    pptx = build_pptx(data, DELIVERABLES / "AFRi_Marigold_Initial_Concepts.pptx")
    report = {"ok": pdf["bytes"] > 0, **pdf, **pptx,
              "shots_per_concept": {e["concept"]["name"]: len(e["shots"])
                                    for e in data}}
    print(json.dumps(report, indent=1))
    return 0 if report["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
