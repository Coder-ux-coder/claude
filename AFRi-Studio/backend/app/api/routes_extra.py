"""Refinement, export and reference-import routes."""
from __future__ import annotations

import json
import shutil
import zipfile
from pathlib import Path

from fastapi import APIRouter, Body, File, HTTPException, UploadFile
from pydantic import BaseModel

from backend.app.agents import refine
from backend.app.core.config import DELIVERABLES, PROJECTS_DIR, ROOT
from backend.app.core.security import UnsafePath, safe_asset_path
from backend.app.database.db import get_conn, new_id, now, tx
from backend.app.versions import store

router = APIRouter(prefix="/api")


# ------------------------------------------------------------- refinement
class RefineIn(BaseModel):
    concept_id: str
    max_iterations: int = 5
    max_seconds: int = 900
    render_each: bool = False


@router.post("/refine/start")
def refine_start(body: RefineIn):
    try:
        return refine.start_run(body.concept_id, body.max_iterations,
                                body.max_seconds, body.render_each)
    except ValueError as exc:
        raise HTTPException(404, str(exc))


@router.post("/refine/{rid}/stop")
def refine_stop(rid: str):
    r = refine.stop_run(rid)
    if not r:
        raise HTTPException(404, "run not found")
    return r


@router.get("/refine/runs")
def refine_runs(concept_id: str | None = None):
    return {"runs": refine.list_runs(concept_id)}


@router.get("/refine/{rid}")
def refine_get(rid: str):
    r = refine.get_run(rid)
    if not r:
        raise HTTPException(404, "run not found")
    return r


@router.get("/refine/measure/{version_id}")
def refine_measure(version_id: str):
    v = store.get_version(version_id)
    if not v:
        raise HTTPException(404, "version not found")
    from design_engine.configurations.schema import DesignConfig
    return {"measurements": refine.measure(DesignConfig(**v["config"]))}


# ----------------------------------------------------------------- export
class ExportIn(BaseModel):
    version_ids: list[str]
    name: str = "afri_marigold_delivery"


@router.post("/exports")
def build_export(body: ExportIn):
    DELIVERABLES.mkdir(parents=True, exist_ok=True)
    out = DELIVERABLES / f"{body.name}.zip"
    included, missing = [], []
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for vid in body.version_ids:
            v = store.get_version(vid)
            if not v:
                missing.append(vid)
                continue
            concept = store.get_concept(v["concept_id"])
            folder = f"{concept['name'].replace(' ', '_')}_v{v['number']}"
            z.writestr(f"{folder}/configuration.json",
                       json.dumps(v["config"], indent=2))
            z.writestr(f"{folder}/validation.json",
                       json.dumps(v["validation"], indent=2))
            for key, rel in (v["assets"] or {}).items():
                p = ROOT / rel
                if p.is_file():
                    z.write(p, f"{folder}/{p.name}")
                    included.append(f"{folder}/{p.name}")
                else:
                    missing.append(rel)
        readme = ROOT / "deliverables" / "README_DELIVERY.md"
        if readme.is_file():
            z.write(readme, "README.md")
    if not out.is_file() or out.stat().st_size == 0:
        raise HTTPException(500, "export produced no file")
    return {"path": str(out.relative_to(ROOT)), "bytes": out.stat().st_size,
            "included": included, "missing": missing}


@router.get("/exports")
def list_exports():
    if not DELIVERABLES.exists():
        return {"exports": []}
    return {"exports": [
        {"name": p.name, "bytes": p.stat().st_size, "modified": p.stat().st_mtime,
         "path": str(p.relative_to(ROOT))}
        for p in sorted(DELIVERABLES.iterdir()) if p.is_file()]}


@router.get("/exports/download/{name}")
def download_export(name: str):
    from fastapi.responses import FileResponse
    try:
        p = safe_asset_path("deliverables", name)
    except UnsafePath as exc:
        raise HTTPException(400, str(exc))
    if not p.is_file():
        raise HTTPException(404, "export not found")
    return FileResponse(p, filename=name)


# ------------------------------------------------------------- references
MAX_REFERENCE_MB = 40
ALLOWED_REFERENCE = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp",
                     ".pptx", ".pdf", ".svg"}


@router.post("/references/import")
async def import_reference(concept_id: str | None = None,
                           file: UploadFile = File(...)):
    """Import a client reference. Originals are preserved and never uploaded."""
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in ALLOWED_REFERENCE:
        raise HTTPException(
            422, f"unsupported reference type '{suffix}'. Allowed: "
                 f"{', '.join(sorted(ALLOWED_REFERENCE))}")
    data = await file.read()
    if len(data) > MAX_REFERENCE_MB * 1024 * 1024:
        raise HTTPException(413, f"reference exceeds {MAX_REFERENCE_MB} MB")

    refs = PROJECTS_DIR / "afri_marigold" / "references"
    refs.mkdir(parents=True, exist_ok=True)
    rid = new_id("ref")
    safe = "".join(ch for ch in (file.filename or "reference")
                   if ch.isalnum() or ch in "._- ")[:80]
    dest = refs / f"{rid}_{safe}"
    dest.write_bytes(data)

    extracted = []
    if suffix == ".pptx":
        extracted = _extract_pptx(dest, refs / rid)

    with tx() as c:
        c.execute("INSERT INTO references_ (id,concept_id,filename,kind,path,"
                  "notes,created_at) VALUES (?,?,?,?,?,?,?)",
                  (rid, concept_id, safe, suffix.lstrip("."),
                   str(dest.relative_to(ROOT)),
                   f"{len(extracted)} embedded items extracted" if extracted else "",
                   now()))
    return {"id": rid, "filename": safe, "bytes": len(data),
            "extracted": extracted,
            "note": "Stored locally only. Importing a reference never overwrites "
                    "existing geometry or approved versions."}


def _extract_pptx(path: Path, outdir: Path) -> list[str]:
    """Pull images and text out of a PowerPoint using local tools only."""
    outdir.mkdir(parents=True, exist_ok=True)
    found = []
    try:
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.startswith("ppt/media/"):
                    target = outdir / Path(n).name
                    target.write_bytes(z.read(n))
                    found.append(str(target.relative_to(ROOT)))
    except Exception:
        return found
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        lines = []
        for i, slide in enumerate(prs.slides, 1):
            lines.append(f"## Slide {i}")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    lines.append(shape.text_frame.text.strip())
        txt = outdir / "extracted_text.md"
        txt.write_text("\n\n".join(lines))
        found.append(str(txt.relative_to(ROOT)))
    except Exception:
        pass
    return found


@router.get("/references")
def list_references(concept_id: str | None = None):
    q = "SELECT * FROM references_"
    args: list = []
    if concept_id:
        q += " WHERE concept_id=?"
        args.append(concept_id)
    q += " ORDER BY created_at DESC"
    return {"references": [dict(r) for r in get_conn().execute(q, args).fetchall()]}
